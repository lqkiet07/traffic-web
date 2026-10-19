import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_appendix_docx():
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(18)
    title_run = title_p.add_run("PHỤ LỤC: HÌNH ẢNH GIAO DIỆN HỆ THỐNG TRAFFICDT-WEB")
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    figures = [
        (
            "docs/figures/hinh_1_replay_map.png",
            "Hình 1. Giao diện bản đồ tái hiện mô phỏng mạng lưới giao thông vi mô 8 nút giao",
            6.0
        ),
        (
            "docs/figures/hinh_2_inspector.png",
            "Hình 2. Giao diện chi tiết chu kỳ đèn và cột tín hiệu tại nút giao",
            6.0
        ),
        (
            "docs/figures/hinh_4_heterogeneous.png",
            "Hình 3. Giao diện 3D trực quan hóa thuật toán bằng dòng giao thông",
            6.0
        ),
        (
            "docs/figures/hinh_5_stepper_math.png",
            "Hình 4. Giao diện phân bổ thời lượng xanh",
            6.0
        ),
        (
            "docs/figures/hinh_6_camera_yolo.png",
            "Hình 5. Giao diện nhận diện phương tiện YOLOv8 và vùng quan tâm ROI trên luồng camera",
            5.8
        ),
    ]

    for img_path, caption_text, width_in in figures:
        if not os.path.exists(img_path):
            print(f"⚠️ Không tìm thấy ảnh: {img_path}")
            continue

        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(20)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(12)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(30, 41, 59)

    out_path = "docs/PHU_LUC_HINH_ANH.docx"
    doc.save(out_path)
    file_size_mb = os.path.getsize(out_path) / (1024 * 1024)
    print(f"✅ Đã tạo thành công file Word tại: {out_path} ({file_size_mb:.2f} MB)")

if __name__ == "__main__":
    create_appendix_docx()
